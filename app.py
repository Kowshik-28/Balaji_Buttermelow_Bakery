import json
import os
import io
import re
import shutil
import uuid
from contextlib import asynccontextmanager
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook

from dotenv import load_dotenv
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware

from auth import hash_password, session_secret, verify_password
from database import get_db, initialize_database, rows_to_dicts, utc_now
from notifications import send_order_notifications


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

ADMIN_USERNAME = os.getenv("ADMIN_USERNAME", "owner")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "Sweety@1234")
ORDER_STATUSES = ("new", "confirmed", "preparing",
                  "ready", "completed", "cancelled")
ALLOWED_IMAGE_TYPES = {"image/jpeg": ".jpg",
                       "image/png": ".png", "image/webp": ".webp"}


@asynccontextmanager
async def lifespan(_app: FastAPI):
    initialize_database(hash_password(ADMIN_PASSWORD), ADMIN_USERNAME)
    yield


app = FastAPI(title="Buttermelow", lifespan=lifespan)
app.add_middleware(
    SessionMiddleware,
    secret_key=session_secret(),
    same_site="lax",
    https_only=os.getenv("COOKIE_SECURE", "false").lower() == "true",
)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
app.mount("/uploads", StaticFiles(directory=BASE_DIR / "uploads"), name="uploads")
templates = Jinja2Templates(directory=BASE_DIR / "templates")


def money(paise: int) -> str:
    return f"₹{paise / 100:,.2f}"


templates.env.filters["money"] = money


def admin_required(request: Request):
    if not request.session.get("admin_id"):
        return RedirectResponse("/owner/login", status_code=303)
    return None


def normalize_phone(phone: str) -> str:
    value = re.sub(r"[^\d+]", "", phone.strip())
    if value.startswith("00"):
        value = f"+{value[2:]}"
    if not value.startswith("+") and len(value) == 10:
        value = f"+91{value}"
    if not re.fullmatch(r"\+\d{10,15}", value):
        raise ValueError("Enter a valid phone number, including country code.")
    return value


def menu_items(available_only=True):
    query = "SELECT * FROM menu_items"
    if available_only:
        query += " WHERE is_available = 1"
    query += " ORDER BY category, name"
    with get_db() as db:
        return rows_to_dicts(db.execute(query).fetchall())


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    items = menu_items(available_only=False)[:4]
    return templates.TemplateResponse(
        request, "home.html", {"items": items, "page": "home"}
    )


@app.get("/health")
def health_check():
    return {"status": "healthy", "timestamp": utc_now()}


@app.get("/menu", response_class=HTMLResponse)
def menu(request: Request):
    return templates.TemplateResponse(
        request, "menu.html", {"items": menu_items(available_only=False), "page": "menu"}
    )


@app.get("/cart", response_class=HTMLResponse)
def cart(request: Request):
    return templates.TemplateResponse(request, "cart.html", {"page": "cart"})


@app.get("/contact", response_class=HTMLResponse)
def contact(request: Request):
    return templates.TemplateResponse(request, "contact.html", {"page": "contact"})


@app.get("/api/menu")
def api_menu():
    return menu_items()


@app.post("/api/orders")
async def create_order(request: Request, background_tasks: BackgroundTasks):
    payload = await request.json()
    raw_items = payload.get("items") or []
    if not raw_items:
        raise HTTPException(400, "Your cart is empty.")

    try:
        customer_name = str(payload["customer"]["name"]).strip()
        customer_phone = normalize_phone(str(payload["customer"]["phone"]))
    except (KeyError, TypeError, ValueError) as exc:
        raise HTTPException(
            400, str(exc) or "Customer name and phone are required.")
    if len(customer_name) < 2:
        raise HTTPException(400, "Enter the customer's full name.")

    email = str(payload["customer"].get("email", "")).strip() or None
    address = str(payload["customer"].get("address", "")).strip() or None
    fulfillment = str(payload.get("fulfillment_type", "pickup"))
    if fulfillment not in {"pickup", "delivery"}:
        fulfillment = "pickup"
    if fulfillment == "delivery" and not address:
        raise HTTPException(400, "Delivery address is required.")

    quantities = {}
    for item in raw_items:
        try:
            item_id = int(item["id"])
            quantity = int(item["quantity"])
        except (KeyError, TypeError, ValueError):
            raise HTTPException(400, "Invalid cart item.")
        if quantity < 1 or quantity > 50:
            raise HTTPException(400, "Item quantity must be between 1 and 50.")
        quantities[item_id] = quantities.get(item_id, 0) + quantity

    placeholders = ",".join("?" for _ in quantities)
    now = utc_now()
    with get_db() as db:
        products = db.execute(
            f"""
            SELECT * FROM menu_items
            WHERE id IN ({placeholders}) AND is_available = 1
            """,
            tuple(quantities),
        ).fetchall()
        if len(products) != len(quantities):
            raise HTTPException(
                400, "One or more cart items are no longer available.")

        customer = db.execute(
            "SELECT * FROM customers WHERE phone = ?", (customer_phone,)
        ).fetchone()
        if customer:
            db.execute(
                """
                UPDATE customers
                SET name = ?, email = COALESCE(?, email),
                    address = COALESCE(?, address),
                    order_count = order_count + 1, last_seen_at = ?
                WHERE id = ?
                """,
                (customer_name, email, address, now, customer["id"]),
            )
            customer_id = customer["id"]
            returning_customer = True
        else:
            cursor = db.execute(
                """
                INSERT INTO customers
                    (name, phone, email, address, order_count, first_seen_at, last_seen_at)
                VALUES (?, ?, ?, ?, 1, ?, ?)
                """,
                (customer_name, customer_phone, email, address, now, now),
            )
            customer_id = cursor.lastrowid
            returning_customer = False

        total = sum(product["price_paise"] * quantities[product["id"]]
                    for product in products)
        cursor = db.execute(
            """
            INSERT INTO orders
                (customer_id, fulfillment_type, requested_time, notes,
                 total_paise, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                customer_id,
                fulfillment,
                str(payload.get("requested_time", "")).strip() or None,
                str(payload.get("notes", "")).strip() or None,
                total,
                now,
                now,
            ),
        )
        order_id = cursor.lastrowid
        order_number = f"TBM-{order_id:05d}"
        db.execute(
            "UPDATE orders SET order_number = ? WHERE id = ?",
            (order_number, order_id),
        )
        order_items = []
        for product in products:
            quantity = quantities[product["id"]]
            subtotal = product["price_paise"] * quantity
            db.execute(
                """
                INSERT INTO order_items
                    (order_id, menu_item_id, item_name, quantity,
                     unit_price_paise, subtotal_paise)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    order_id,
                    product["id"],
                    product["name"],
                    quantity,
                    product["price_paise"],
                    subtotal,
                ),
            )
            order_items.append(
                {
                    "item_name": product["name"],
                    "quantity": quantity,
                    "unit_price_paise": product["price_paise"],
                }
            )

    notification_order = {
        "id": order_id,
        "order_number": order_number,
        "customer_name": customer_name,
        "customer_phone": customer_phone,
        "total_paise": total,
    }
    background_tasks.add_task(
        send_order_notifications, notification_order, order_items
    )
    return {
        "ok": True,
        "order_number": order_number,
        "returning_customer": returning_customer,
    }


@app.get("/order/success/{order_number}", response_class=HTMLResponse)
def order_success(request: Request, order_number: str):
    with get_db() as db:
        order = db.execute(
            """
            SELECT o.*, c.name AS customer_name
            FROM orders o JOIN customers c ON c.id = o.customer_id
            WHERE o.order_number = ?
            """,
            (order_number,),
        ).fetchone()
    if not order:
        raise HTTPException(404)
    return templates.TemplateResponse(
        request, "order_success.html", {"order": dict(order), "page": ""}
    )


@app.get("/owner/login", response_class=HTMLResponse)
def owner_login(request: Request):
    if request.session.get("admin_id"):
        return RedirectResponse("/owner", status_code=303)
    return templates.TemplateResponse(request, "owner_login.html", {"error": None})


@app.post("/owner/login", response_class=HTMLResponse)
def owner_login_submit(
    request: Request, username: str = Form(...), password: str = Form(...)
):
    with get_db() as db:
        admin = db.execute(
            "SELECT * FROM admins WHERE username = ?", (username.strip(),)
        ).fetchone()
    if not admin or not verify_password(password, admin["password_hash"]):
        return templates.TemplateResponse(
            request,
            "owner_login.html",
            {"error": "Incorrect username or password."},
            status_code=401,
        )
    request.session.clear()
    request.session["admin_id"] = admin["id"]
    request.session["admin_username"] = admin["username"]
    return RedirectResponse("/owner", status_code=303)


@app.post("/owner/logout")
def owner_logout(request: Request):
    request.session.clear()
    return RedirectResponse("/owner/login", status_code=303)


@app.get("/owner", response_class=HTMLResponse)
def owner_dashboard(request: Request, archived: int = 0):
    redirect = admin_required(request)
    if redirect:
        return redirect
    with get_db() as db:
        orders = rows_to_dicts(
            db.execute(
                """
                SELECT o.*, c.name AS customer_name, c.phone AS customer_phone,
                       c.order_count AS customer_order_count, c.address AS customer_address
                FROM orders o JOIN customers c ON c.id = o.customer_id
                WHERE o.is_archived = ?
                ORDER BY o.created_at DESC
                """,
                (1 if archived else 0,),
            ).fetchall()
        )
        for order in orders:
            order["items"] = rows_to_dicts(
                db.execute(
                    "SELECT * FROM order_items WHERE order_id = ? ORDER BY id",
                    (order["id"],),
                ).fetchall()
            )
        stats = dict(
            db.execute(
                """
                SELECT
                    COUNT(*) AS total_orders,
                    COALESCE(SUM(CASE WHEN status = 'new' AND is_archived = 0 THEN 1 ELSE 0 END), 0) AS new_orders,
                    COALESCE(SUM(CASE WHEN status = 'completed' THEN total_paise ELSE 0 END), 0) AS completed_revenue
                FROM orders
                """
            ).fetchone()
        )
        customer_count = db.execute(
            "SELECT COUNT(*) FROM customers").fetchone()[0]
    stats["customer_count"] = customer_count
    return templates.TemplateResponse(
        request,
        "owner_dashboard.html",
        {
            "orders": orders,
            "stats": stats,
            "statuses": ORDER_STATUSES,
            "archived": bool(archived),
            "admin_username": request.session["admin_username"],
        },
    )


@app.post("/owner/orders/{order_id}/status")
def update_order_status(
    request: Request, order_id: int, status: str = Form(...)
):
    redirect = admin_required(request)
    if redirect:
        return redirect
    if status not in ORDER_STATUSES:
        raise HTTPException(400, "Unknown order status.")
    with get_db() as db:
        db.execute(
            "UPDATE orders SET status = ?, updated_at = ? WHERE id = ?",
            (status, utc_now(), order_id),
        )
    return RedirectResponse("/owner", status_code=303)


@app.post("/owner/orders/{order_id}/archive")
def archive_order(request: Request, order_id: int):
    redirect = admin_required(request)
    if redirect:
        return redirect
    with get_db() as db:
        order = db.execute(
            "SELECT status FROM orders WHERE id = ?", (order_id,)).fetchone()
        if not order:
            raise HTTPException(404)
        if order["status"] not in {"completed", "cancelled"}:
            raise HTTPException(
                400, "Complete or cancel the order before archiving it.")
        db.execute(
            "UPDATE orders SET is_archived = 1, updated_at = ? WHERE id = ?",
            (utc_now(), order_id),
        )
    return RedirectResponse("/owner", status_code=303)


@app.post("/owner/orders/{order_id}/delete")
def delete_order(request: Request, order_id: int):
    redirect = admin_required(request)
    if redirect:
        return redirect
    with get_db() as db:
        order = db.execute(
            "SELECT customer_id, status FROM orders WHERE id = ?", (order_id,)
        ).fetchone()
        if not order:
            raise HTTPException(404, "Order not found.")
        if order["status"] != "cancelled":
            raise HTTPException(
                400, "Only cancelled orders can be deleted."
            )
        customer_id = order["customer_id"]
        
        # Cascading deletes will clean up order_items and notifications
        db.execute("DELETE FROM orders WHERE id = ?", (order_id,))
        
        # Decrement customer's order count by 1 (using max() to prevent negative counts)
        db.execute(
            """
            UPDATE customers
            SET order_count = max(0, order_count - 1)
            WHERE id = ?
            """,
            (customer_id,),
        )
    return RedirectResponse("/owner", status_code=303)


@app.get("/owner/menu", response_class=HTMLResponse)
def owner_menu(request: Request):
    redirect = admin_required(request)
    if redirect:
        return redirect
    return templates.TemplateResponse(
        request,
        "owner_menu.html",
        {"items": menu_items(
            False), "admin_username": request.session["admin_username"]},
    )


async def save_upload(image: UploadFile | None) -> str | None:
    if not image or not image.filename:
        return None
    extension = ALLOWED_IMAGE_TYPES.get(image.content_type or "")
    if not extension:
        raise HTTPException(400, "Upload a JPG, PNG, or WebP image.")
    target_name = f"{uuid.uuid4().hex}{extension}"
    target = BASE_DIR / "uploads" / target_name
    with target.open("wb") as destination:
        shutil.copyfileobj(image.file, destination)
    if target.stat().st_size > 5 * 1024 * 1024:
        target.unlink(missing_ok=True)
        raise HTTPException(400, "Image must be smaller than 5 MB.")
    return f"uploads/{target_name}"


@app.post("/owner/menu")
async def add_menu_item(
    request: Request,
    name: str = Form(...),
    category: str = Form(...),
    description: str = Form(""),
    price: float = Form(...),
    image: UploadFile | None = File(None),
):
    redirect = admin_required(request)
    if redirect:
        return redirect
    image_path = await save_upload(image)
    now = utc_now()
    with get_db() as db:
        db.execute(
            """
            INSERT INTO menu_items
                (name, category, description, price_paise, image_path, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                name.strip(),
                category.strip(),
                description.strip(),
                max(0, round(price * 100)),
                image_path,
                now,
                now,
            ),
        )
    return RedirectResponse("/owner/menu", status_code=303)


@app.post("/owner/menu/{item_id}")
async def update_menu_item(
    request: Request,
    item_id: int,
    name: str = Form(...),
    category: str = Form(...),
    description: str = Form(""),
    price: float = Form(...),
    is_available: str | None = Form(None),
    image: UploadFile | None = File(None),
):
    redirect = admin_required(request)
    if redirect:
        return redirect
    image_path = await save_upload(image)
    with get_db() as db:
        current = db.execute(
            "SELECT image_path FROM menu_items WHERE id = ?", (item_id,)
        ).fetchone()
        if not current:
            raise HTTPException(404)
        db.execute(
            """
            UPDATE menu_items
            SET name = ?, category = ?, description = ?, price_paise = ?,
                image_path = ?, is_available = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                name.strip(),
                category.strip(),
                description.strip(),
                max(0, round(price * 100)),
                image_path or current["image_path"],
                1 if is_available else 0,
                utc_now(),
                item_id,
            ),
        )
    return RedirectResponse("/owner/menu", status_code=303)


@app.post("/owner/menu/{item_id}/delete")
def delete_menu_item(request: Request, item_id: int):
    redirect = admin_required(request)
    if redirect:
        return redirect
    with get_db() as db:
        item = db.execute("SELECT image_path FROM menu_items WHERE id = ?", (item_id,)).fetchone()
        if not item:
            raise HTTPException(404, "Menu item not found")
        if item["image_path"] and item["image_path"].startswith("uploads/"):
            img_path = BASE_DIR / item["image_path"]
            img_path.unlink(missing_ok=True)
        db.execute("DELETE FROM menu_items WHERE id = ?", (item_id,))
    return RedirectResponse("/owner/menu", status_code=303)



@app.get("/owner/customers", response_class=HTMLResponse)
def owner_customers(request: Request):
    redirect = admin_required(request)
    if redirect:
        return redirect
    with get_db() as db:
        customers = rows_to_dicts(
            db.execute(
                "SELECT * FROM customers ORDER BY last_seen_at DESC"
            ).fetchall()
        )
    return templates.TemplateResponse(
        request,
        "owner_customers.html",
        {"customers": customers,
            "admin_username": request.session["admin_username"]},
    )


@app.exception_handler(HTTPException)
async def http_exception_handler(request: Request, exc: HTTPException):
    if request.url.path.startswith("/api/"):
        return JSONResponse({"detail": exc.detail}, status_code=exc.status_code)
    return templates.TemplateResponse(
        request,
        "error.html",
        {"message": exc.detail or "Something went wrong."},
        status_code=exc.status_code,
    )

@app.get("/owner/export-archived")
def export_archived_orders(request: Request):
    redirect = admin_required(request)
    if redirect:
        return redirect

    with get_db() as db:
        rows = db.execute(
            """
            SELECT o.*, c.name AS customer_name, c.phone AS customer_phone,
                   c.address AS customer_address
            FROM orders o JOIN customers c ON c.id = o.customer_id
            WHERE o.is_archived = 1
            ORDER BY o.created_at
            """
        ).fetchall()
        orders = rows_to_dicts(rows)
        if not orders:
            raise HTTPException(status_code=400, detail="No archived orders found to export.")
        for order in orders:
            order["items"] = rows_to_dicts(
                db.execute(
                    "SELECT * FROM order_items WHERE order_id = ? ORDER BY id",
                    (order["id"],),
                ).fetchall()
            )

    # Group by created date (YYYY-MM-DD) and format sheet name like "15 JUN 2026"
    groups = {}
    for o in orders:
        created = o.get("created_at") or ""
        date_part = ""
        try:
            # try ISO or space-separated timestamp
            dt = datetime.fromisoformat(created.replace("Z", "")) if created else None
            if dt:
                date_part = dt.strftime("%d %b %Y").upper()
        except Exception:
            # fallback: take first token YYYY-MM-DD or raw
            date_part = created.split("T")[0] if "T" in created else (created.split(" ")[0] if created else "UNKNOWN")
            try:
                dt = datetime.fromisoformat(date_part)
                date_part = dt.strftime("%d %b %Y").upper()
            except Exception:
                date_part = (date_part or "UNKNOWN").upper()
        groups.setdefault(date_part or "UNKNOWN", []).append(o)

    wb = Workbook()
    # remove default sheet
    default = wb.active
    wb.remove(default)

    for date_key, ods in sorted(groups.items()):
        # Clean invalid characters for Excel sheet name (\, /, ?, *, :, [, ])
        clean_date_key = re.sub(r"[\\*?:/\[\]]", "_", date_key)
        sheet_name = clean_date_key[:31]  # Excel sheet name limit
        ws = wb.create_sheet(title=sheet_name)
        ws.append(["TBM", "Customer name", "Items", "Phone no", "Address", "Amount", "Time of delivery"])
        for o in ods:
            items_text = "; ".join(
                f"{it['quantity']}× {it['item_name']}" for it in o.get("items", [])
            )
            amount = money(o.get("total_paise", 0))
            delivery_time = o.get("requested_time") or ""
            # normalize delivery_time display
            if isinstance(delivery_time, str):
                delivery_time = delivery_time.replace("T", " ")
            ws.append(
                [
                    o.get("order_number"),
                    o.get("customer_name") or "",
                    items_text,
                    o.get("customer_phone") or "",
                    o.get("customer_address") or "",
                    amount,
                    delivery_time,
                ]
            )

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"archived-orders-{utc_now().split('T')[0]}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )